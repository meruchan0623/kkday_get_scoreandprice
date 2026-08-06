from dataclasses import replace
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

import kkday_scraper.exporter as exporter_module
from kkday_scraper.exporter import RunExporter
from kkday_scraper.models import (
    FailureKind,
    FailureRecord,
    InputItem,
    PriceStatus,
    ProductResult,
    ProductSnapshot,
    RunConfig,
    SkuRow,
)


def make_product() -> ProductResult:
    item = InputItem(
        "https://www.kkday.com/zh-tw/product/149025-example", vendor="TSIM 網址"
    )
    row = SkuRow(
        vendor=item.vendor,
        input_order=0,
        product_id="149025",
        product_url=item.url,
        product_name="日本 eSIM",
        package_id="476215",
        package_name="每日/无限流量",
        item_id="221992",
        sku_id="sku-1",
        usage_date=date(2026, 7, 22),
        options=("500MB/日", "1 日"),
        available=True,
        currency="KRW",
        sale_price=Decimal("711"),
        official_price=None,
        price_source="single_price.fullday",
        price_status=PriceStatus.TRUSTED,
    )
    failure = FailureRecord(
        FailureKind.API, item.url, "package:2", "HTTP 500", True
    )
    snapshot = ProductSnapshot(
        product_id="149025",
        url=item.url,
        name="日本 eSIM",
        packages=(),
        destination="日本",
        rating=4.75,
        review_count=241,
        sold_count=2112,
        wishlist_count=2061,
        page_min_price=Decimal("55"),
        confirmation="立即确认",
        cancellation_policy="不可取消",
        features=("日本当地信号", "iPhone XR 以上设备"),
    )
    return ProductResult(item, snapshot, (row,), (failure,), ("一个警告",))


def test_checkpoint_writes_stable_workbook_and_failure_ledger(
    tmp_path: Path,
) -> None:
    config = RunConfig(date(2026, 7, 22), tmp_path)
    exporter = RunExporter(config, run_id="20260722_120000")

    output = exporter.checkpoint((make_product(),), final=True)

    assert output.exists()
    assert output.name == "kkday_price_results.partial.xlsx"
    workbook = openpyxl.load_workbook(output, data_only=True)
    assert workbook.sheetnames == ["抓取结果", "失败商品", "运行汇总"]
    sheet = workbook["抓取结果"]
    headers = [cell.value for cell in sheet[1]]
    assert headers[:8] == [
        "厂商",
        "输入序号",
        "商品ID",
        "URL",
        "商品名称",
        "目的地",
        "评分",
        "评论数",
    ]
    assert sheet["F2"].value == "日本"
    assert sheet["G2"].value == 4.75
    assert sheet["H2"].value == 241
    assert sheet["I2"].value == 2112
    assert sheet["J2"].value == 2061
    assert sheet["K2"].value == 55
    assert sheet["L2"].value == "立即确认"
    assert sheet["M2"].value == "不可取消"
    assert sheet["N2"].value == "日本当地信号 | iPhone XR 以上设备"
    assert sheet["V2"].value == 711
    assert exporter.summary_path.exists()
    assert exporter.failures_path.exists()


def test_exact_result_directory_does_not_add_a_timestamp_folder(
    tmp_path: Path,
) -> None:
    result_dir = tmp_path / "chosen-result"
    exporter = RunExporter(
        RunConfig(date(2026, 7, 22), tmp_path),
        result_dir=result_dir,
    )

    output = exporter.checkpoint((make_product(),), final=True)

    assert output == result_dir / "kkday_price_results.partial.xlsx"
    assert exporter.summary_path == result_dir / "run_summary.json"


def test_default_result_directories_are_unique_within_the_same_second(
    tmp_path: Path,
) -> None:
    config = RunConfig(date(2026, 7, 22), tmp_path)

    first = RunExporter(config)
    second = RunExporter(config)

    assert first.run_dir != second.run_dir
    assert first.batch_id != second.batch_id


def test_default_result_directory_rejects_an_existing_run_id(tmp_path: Path) -> None:
    config = RunConfig(date(2026, 7, 22), tmp_path)
    RunExporter(config, run_id="20260722_120010")

    with pytest.raises(FileExistsError, match="结果目录已存在"):
        RunExporter(config, run_id="20260722_120010")


def test_explicit_result_directory_has_one_writer_and_recovers_after_close(
    tmp_path: Path,
) -> None:
    result_dir = tmp_path / "chosen-result"
    first = RunExporter(RunConfig(date(2026, 7, 22), tmp_path), result_dir=result_dir)

    first.checkpoint((make_product(),), final=False)
    first.checkpoint((make_product(),), final=True)
    with pytest.raises(RuntimeError, match="正在被另一个导出任务占用"):
        RunExporter(RunConfig(date(2026, 7, 22), tmp_path), result_dir=result_dir)

    first.close()
    recovered = RunExporter(RunConfig(date(2026, 7, 22), tmp_path), result_dir=result_dir)
    try:
        assert recovered.checkpoint((make_product(),), final=True).exists()
    finally:
        recovered.close()


def test_explicit_result_directory_lock_has_a_non_posix_fallback(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(exporter_module, "fcntl", None)
    result_dir = tmp_path / "portable-result"
    first = RunExporter(RunConfig(date(2026, 7, 22), tmp_path), result_dir=result_dir)

    with pytest.raises(RuntimeError, match="正在被另一个导出任务占用"):
        RunExporter(RunConfig(date(2026, 7, 22), tmp_path), result_dir=result_dir)

    first.close()
    recovered = RunExporter(RunConfig(date(2026, 7, 22), tmp_path), result_dir=result_dir)
    recovered.close()


def test_checkpoint_writes_one_batch_id_to_json_and_workbook(tmp_path: Path) -> None:
    exporter = RunExporter(
        RunConfig(date(2026, 7, 22), tmp_path), run_id="20260722_120011"
    )

    output = exporter.checkpoint((make_product(),), final=True)

    summary = __import__("json").loads(exporter.summary_path.read_text(encoding="utf-8"))
    sheet = openpyxl.load_workbook(output, data_only=True)["运行汇总"]
    workbook_summary = {row[0].value: row[1].value for row in sheet.iter_rows(min_row=2)}
    assert summary["batch_id"] == exporter.batch_id
    assert workbook_summary["批次ID"] == exporter.batch_id


def test_checkpoint_omits_unavailable_rows_by_default(tmp_path: Path) -> None:
    product = make_product()
    unavailable = product.rows[0].__class__(
        **{
            **product.rows[0].__dict__,
            "sku_id": "closed",
            "available": False,
            "price_status": PriceStatus.UNAVAILABLE,
        }
    )
    product = ProductResult(product.input_item, None, product.rows + (unavailable,))
    exporter = RunExporter(
        RunConfig(date(2026, 7, 22), tmp_path), run_id="20260722_120001"
    )

    output = exporter.checkpoint((product,), final=False)

    workbook = openpyxl.load_workbook(output, data_only=True)
    assert workbook["抓取结果"].max_row == 2


def test_failed_checkpoint_preserves_previous_valid_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    exporter = RunExporter(
        RunConfig(date(2026, 7, 22), tmp_path), run_id="20260722_120002"
    )
    output = exporter.checkpoint((make_product(),), final=False)
    previous = output.read_bytes()

    def fail_save(*_args, **_kwargs):
        raise OSError("disk full")

    monkeypatch.setattr(openpyxl.Workbook, "save", fail_save)
    with pytest.raises(OSError, match="disk full"):
        exporter.checkpoint((make_product(),), final=True)

    assert output.read_bytes() == previous


def test_remote_text_cannot_become_an_excel_formula(tmp_path: Path) -> None:
    product = make_product()
    dangerous = replace(
        product.rows[0],
        product_name='=HYPERLINK("https://example.invalid")',
        package_name="@SUM(1,1)",
    )
    product = ProductResult(product.input_item, None, (dangerous,))
    exporter = RunExporter(
        RunConfig(date(2026, 7, 22), tmp_path), run_id="20260722_120003"
    )

    output = exporter.checkpoint((product,), final=True)

    workbook = openpyxl.load_workbook(output, data_only=False)
    assert workbook["抓取结果"]["E2"].value.startswith("'=")
    assert workbook["抓取结果"]["P2"].value.startswith("'@")


def test_checkpoint_sorts_tree_order_and_keeps_product_fields_on_first_row_only(
    tmp_path: Path,
) -> None:
    product = make_product()
    later = replace(product.rows[0], sku_id="later", tree_order=(1, 0, 1, 0))
    first = replace(product.rows[0], sku_id="first", tree_order=(0, 0, 0, 0))
    product = replace(product, rows=(later, first))
    exporter = RunExporter(
        RunConfig(date(2026, 7, 22), tmp_path), run_id="20260722_120004"
    )

    output = exporter.checkpoint((product,), final=True)

    sheet = openpyxl.load_workbook(output, data_only=True)["抓取结果"]
    assert [sheet["R2"].value, sheet["R3"].value] == ["first", "later"]
    assert sheet["A2"].value == "TSIM 網址"
    assert sheet["E2"].value == "日本 eSIM"
    assert sheet["F2"].value == "日本"
    assert sheet["A3"].value is None
    assert sheet["E3"].value is None
    assert sheet["F3"].value is None


def test_partial_checkpoint_has_explicit_name_and_covers_zero_sku_url(tmp_path: Path) -> None:
    config = RunConfig(
        date(2026, 7, 22),
        tmp_path,
        requested_count=20,
        ego_task_space="kkday-20-products",
    )
    exporter = RunExporter(config)
    product = ProductResult(
        InputItem("https://www.kkday.com/zh-tw/product/121004-example"),
        None,
        (),
    )

    output = exporter.checkpoint((product,), final=False)

    assert output.name == "kkday_price_results.partial.xlsx"
    workbook = openpyxl.load_workbook(output, read_only=True)
    try:
        rows = list(workbook["失败商品"].iter_rows(values_only=True))
        summary = dict(list(workbook["运行汇总"].iter_rows(values_only=True))[1:])
    finally:
        workbook.close()
    assert any(
        row[0] == "no_available_sku" and row[1] == product.input_item.url
        for row in rows
    )
    assert summary["输入商品数"] == 20
    assert summary["已处理商品数"] == 1
    assert summary["Task space"] == "kkday-20-products"
