from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
import pytest

from kkday_scraper.inputs import (
    deduplicate_inputs,
    load_manual_urls,
    load_txt,
    load_vendor_excel,
    market_from_url,
    normalize_product_url,
)


URL_A = "https://www.kkday.com/zh-tw/product/149025-example"
URL_B = "https://www.kkday.com/zh-tw/product/1982144-example"
MAX_INPUT_FILE_BYTES = 5 * 1024 * 1024
MAX_INPUT_ROWS = 10_000
MAX_INPUT_COLUMNS = 100
MAX_INPUT_CELL_CHARS = 4_096
MAX_INPUT_URL_CHARS = 2_048
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024


def test_normalize_product_url_rejects_non_product_url() -> None:
    with pytest.raises(ValueError, match="KKday 商品 URL"):
        normalize_product_url("https://www.kkday.com/zh-tw/portable-travel-wifi-sim-card")


def test_normalize_product_url_requires_a_numeric_product_id() -> None:
    with pytest.raises(ValueError, match="KKday 商品 URL"):
        normalize_product_url("https://www.kkday.com/zh-tw/product/not-a-number-example")


def test_normalize_product_url_limits_url_length() -> None:
    raw = URL_A + "x" * (MAX_INPUT_URL_CHARS - len(URL_A) + 1)

    with pytest.raises(ValueError, match="URL 长度超过限制"):
        normalize_product_url(raw)


def test_normalize_product_url_drops_query_tokens_and_fragments() -> None:
    raw = f"{URL_A}?previewToken=secret&utm_source=test#package"

    assert normalize_product_url(raw) == URL_A


def test_market_from_url_returns_the_url_market_segment() -> None:
    assert market_from_url("https://www.kkday.com/zh-hk/product/149025-example") == "zh-hk"
    assert market_from_url("https://www.kkday.com/zh-tw/product/149025-example") == "zh-tw"
    assert market_from_url("https://www.kkday.com/product/149025-example") is None


def test_manual_urls_capture_the_market_from_each_url() -> None:
    items = load_manual_urls(
        "https://www.kkday.com/zh-hk/product/149025-example\n"
        "https://www.kkday.com/zh-tw/product/1982144-example\n"
    )
    assert [item.market for item in items] == ["zh-hk", "zh-tw"]


def test_manual_and_txt_inputs_preserve_order(tmp_path: Path) -> None:
    text_path = tmp_path / "urls.txt"
    text_path.write_text(f"{URL_B}\n\n", encoding="utf-8")

    items = load_manual_urls(f"{URL_A}\n{URL_A}") + load_txt(text_path, start_order=2)
    unique = deduplicate_inputs(items)

    assert [item.url for item in unique] == [URL_A, URL_B]
    assert [item.input_order for item in unique] == [0, 2]


def test_manual_urls_enforce_the_same_row_limit_as_files() -> None:
    raw = (URL_A + "\n") * (MAX_INPUT_ROWS + 1)

    with pytest.raises(ValueError, match="输入行数超过限制"):
        load_manual_urls(raw)


def test_load_txt_rejects_oversized_file_before_reading(tmp_path: Path) -> None:
    path = tmp_path / "oversized.txt"
    path.write_bytes(b"x" * (MAX_INPUT_FILE_BYTES + 1))

    with pytest.raises(ValueError, match="输入文件超过大小限制"):
        load_txt(path)


def test_load_txt_enforces_the_row_limit_while_streaming(tmp_path: Path) -> None:
    path = tmp_path / "many-urls.txt"
    path.write_text((URL_A + "\n") * (MAX_INPUT_ROWS + 1), encoding="utf-8")

    with pytest.raises(ValueError, match="输入行数超过限制"):
        load_txt(path)


def test_vendor_excel_reads_headers_and_source_coordinates(tmp_path: Path) -> None:
    path = tmp_path / "vendors.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["TSIM 網址", "BC"])
    sheet.append([URL_A, URL_B])
    sheet.append(["", "not-a-url"])
    workbook.save(path)

    items, warnings = load_vendor_excel(path)

    assert [(item.vendor, item.source_row, item.source_column) for item in items] == [
        ("TSIM 網址", 2, 1),
        ("BC", 2, 2),
    ]
    assert warnings == ["忽略非 KKday 商品 URL：工作表第 3 行，第 2 列：not-a-url"]


def test_vendor_excel_enforces_column_and_cell_limits(tmp_path: Path) -> None:
    too_many_columns = tmp_path / "too-many-columns.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([f"供应商 {index}" for index in range(MAX_INPUT_COLUMNS + 1)])
    workbook.save(too_many_columns)

    with pytest.raises(ValueError, match="输入列数超过限制"):
        load_vendor_excel(too_many_columns)

    oversized_cell = tmp_path / "oversized-cell.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["TSIM 網址"])
    sheet.append(["x" * (MAX_INPUT_CELL_CHARS + 1)])
    workbook.save(oversized_cell)

    with pytest.raises(ValueError, match="单元格内容超过限制"):
        load_vendor_excel(oversized_cell)


def test_vendor_excel_rejects_an_archive_expansion_bomb(tmp_path: Path) -> None:
    path = tmp_path / "expansion-bomb.xlsx"
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr(
            "xl/worksheets/sheet1.xml",
            b"x" * (MAX_XLSX_UNCOMPRESSED_BYTES + 1),
        )

    with pytest.raises(ValueError, match="解压后大小超过限制"):
        load_vendor_excel(path)
