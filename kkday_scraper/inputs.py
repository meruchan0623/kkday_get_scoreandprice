from __future__ import annotations

import io
import re
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit
from zipfile import BadZipFile, ZipFile

import openpyxl

from .models import InputItem, SUPPORTED_MARKETS


MAX_INPUT_FILE_BYTES = 5 * 1024 * 1024
MAX_INPUT_ROWS = 10_000
MAX_INPUT_COLUMNS = 100
MAX_INPUT_CELL_CHARS = 4_096
MAX_INPUT_URL_CHARS = 2_048
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_XLSX_ARCHIVE_ENTRIES = 1_000

PRODUCT_PATH_RE = re.compile(
    r"^/([a-z]{2}-[a-z]{2})/product/(\d+)(?:-|$|/)", re.IGNORECASE
)


def extract_product_id_from_url(raw: str) -> str | None:
    match = PRODUCT_PATH_RE.search(urlsplit(raw.strip()).path)
    return match.group(2) if match else None


def market_from_url(raw: str) -> str | None:
    """Return the KKday market segment (e.g. 'zh-hk') in the URL path, if any."""
    match = PRODUCT_PATH_RE.search(urlsplit(raw.strip()).path)
    if not match:
        return None
    market = match.group(1).lower()
    return market if market in SUPPORTED_MARKETS else None


def _check_file_size(path: Path) -> None:
    if path.stat().st_size > MAX_INPUT_FILE_BYTES:
        raise ValueError(f"输入文件超过大小限制（最多 {MAX_INPUT_FILE_BYTES} 字节）：{path}")


def _check_xlsx_archive(path: Path) -> None:
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
    except (BadZipFile, OSError) as exc:
        raise ValueError(f"Excel 不是有效的 XLSX 文件：{path}") from exc
    if len(entries) > MAX_XLSX_ARCHIVE_ENTRIES:
        raise ValueError(
            f"Excel 压缩包条目超过限制（最多 {MAX_XLSX_ARCHIVE_ENTRIES} 项）"
        )
    expanded_bytes = sum(entry.file_size for entry in entries)
    if expanded_bytes > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise ValueError(
            "Excel 解压后大小超过限制"
            f"（最多 {MAX_XLSX_UNCOMPRESSED_BYTES} 字节）"
        )


def _cell_text(value: object, *, row: int, column: int) -> str:
    text = str(value).strip() if value is not None else ""
    if len(text) > MAX_INPUT_CELL_CHARS:
        raise ValueError(
            "单元格内容超过限制"
            f"（最多 {MAX_INPUT_CELL_CHARS} 字符）：工作表第 {row} 行，第 {column} 列"
        )
    return text


def normalize_product_url(raw: str) -> str:
    value = raw.strip()
    if len(value) > MAX_INPUT_URL_CHARS:
        raise ValueError(f"URL 长度超过限制（最多 {MAX_INPUT_URL_CHARS} 字符）")
    parts = urlsplit(value)
    if parts.scheme not in {"http", "https"} or parts.hostname not in {"kkday.com", "www.kkday.com"}:
        raise ValueError(f"不是有效的 KKday 商品 URL：{raw}")
    if extract_product_id_from_url(value) is None:
        raise ValueError(f"不是有效的 KKday 商品 URL：{raw}")
    return urlunsplit(("https", "www.kkday.com", parts.path, "", ""))


def load_manual_urls(raw: str, start_order: int = 0) -> list[InputItem]:
    items: list[InputItem] = []
    for row_number, line in enumerate(io.StringIO(raw), start=1):
        if row_number > MAX_INPUT_ROWS:
            raise ValueError(f"输入行数超过限制（最多 {MAX_INPUT_ROWS} 行）")
        if not line.strip():
            continue
        items.append(
            InputItem(
                url=normalize_product_url(line),
                input_order=start_order + len(items),
                market=market_from_url(line),
            )
        )
    return items


def load_txt(path: Path, start_order: int = 0) -> list[InputItem]:
    _check_file_size(path)
    items: list[InputItem] = []
    with path.open(encoding="utf-8-sig") as handle:
        for row_number, line in enumerate(handle, start=1):
            if row_number > MAX_INPUT_ROWS:
                raise ValueError(f"输入行数超过限制（最多 {MAX_INPUT_ROWS} 行）：{path}")
            if not line.strip():
                continue
            items.append(
                InputItem(
                    url=normalize_product_url(line),
                    input_order=start_order + len(items),
                    source_path=path,
                    source_row=row_number,
                    market=market_from_url(line),
                )
            )
    return items


def load_vendor_excel(path: Path, start_order: int = 0) -> tuple[list[InputItem], list[str]]:
    _check_file_size(path)
    _check_xlsx_archive(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet.max_row > MAX_INPUT_ROWS + 1:
            raise ValueError(f"输入行数超过限制（最多 {MAX_INPUT_ROWS} 行）")
        if sheet.max_column > MAX_INPUT_COLUMNS:
            raise ValueError(f"输入列数超过限制（最多 {MAX_INPUT_COLUMNS} 列）")
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            raise ValueError("Excel 没有可读取的工作表内容")
        if len(header_row) > MAX_INPUT_COLUMNS:
            raise ValueError(f"输入列数超过限制（最多 {MAX_INPUT_COLUMNS} 列）")
        headers = [
            _cell_text(value, row=1, column=column_index)
            for column_index, value in enumerate(header_row, start=1)
        ]
        items: list[InputItem] = []
        warnings: list[str] = []
        for row_index, row in enumerate(rows, start=2):
            if row_index > MAX_INPUT_ROWS + 1:
                raise ValueError(f"输入行数超过限制（最多 {MAX_INPUT_ROWS} 行）")
            if len(row) > MAX_INPUT_COLUMNS:
                raise ValueError(f"输入列数超过限制（最多 {MAX_INPUT_COLUMNS} 列）")
            for column_index, value in enumerate(row, start=1):
                text = _cell_text(value, row=row_index, column=column_index)
                if not text:
                    continue
                vendor = headers[column_index - 1] if column_index <= len(headers) else ""
                if not vendor:
                    warnings.append(f"忽略无厂商表头的链接：工作表第 {row_index} 行，第 {column_index} 列")
                    continue
                try:
                    url = normalize_product_url(text)
                except ValueError:
                    warnings.append(f"忽略非 KKday 商品 URL：工作表第 {row_index} 行，第 {column_index} 列：{text}")
                    continue
                items.append(
                    InputItem(
                        url=url,
                        vendor=vendor,
                        input_order=start_order + len(items),
                        source_path=path,
                        source_row=row_index,
                        source_column=column_index,
                        market=market_from_url(text),
                    )
                )
        return items, warnings
    finally:
        workbook.close()


def deduplicate_inputs(items: list[InputItem]) -> list[InputItem]:
    seen: set[tuple[str, str]] = set()
    result: list[InputItem] = []
    for item in sorted(items, key=lambda candidate: candidate.input_order):
        key = (item.vendor, item.url)
        if key in seen:
            continue
        seen.add(key)
        result.append(item)
    return result
