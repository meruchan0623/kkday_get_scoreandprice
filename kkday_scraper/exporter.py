from __future__ import annotations

import json
import os
import tempfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

try:
    import fcntl
except ImportError:  # pragma: no cover - exercised via a simulated fallback
    fcntl = None

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .models import ProductResult, ProductSnapshot, RunConfig, SkuRow


BASE_FIELDS = [
    "厂商",
    "输入序号",
    "商品ID",
    "URL",
    "商品名称",
    "目的地",
    "评分",
    "评论数",
    "已售数量",
    "收藏数",
    "页面最低价",
    "确认方式",
    "取消政策",
    "商品卖点",
    "套餐ID",
    "套餐名称",
    "ItemID",
    "SKU ID",
    "使用日期",
    "规格组合",
    "货币",
    "市场",
    "当前价格",
    "原价",
    "价格来源",
    "价格状态",
    "抓取备注",
    "抓取时间",
]


def excel_value(value):
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


class RunExporter:
    def __init__(
        self,
        config: RunConfig,
        run_id: str | None = None,
        result_dir: Path | None = None,
    ) -> None:
        self.config = config
        self.run_id = run_id or (
            datetime.now().strftime("%Y%m%d_%H%M%S_%f")
            + f"_{uuid4().hex[:8]}"
        )
        self.batch_id = self.run_id
        self._lock_fd: int | None = None
        self._lock_path: Path | None = None
        self._uses_marker_lock = False
        if result_dir is None:
            self.run_dir = config.output_dir / self.run_id
            try:
                self.run_dir.mkdir(parents=True, exist_ok=False)
            except FileExistsError as exc:
                raise FileExistsError(f"结果目录已存在，拒绝覆盖：{self.run_dir}") from exc
        else:
            self.run_dir = result_dir
            self.run_dir.mkdir(parents=True, exist_ok=True)
            self._claim_result_dir()
        self.output_path = self.run_dir / "kkday_price_results.xlsx"
        self.partial_output_path = self.run_dir / "kkday_price_results.partial.xlsx"
        self.summary_path = self.run_dir / "run_summary.json"
        self.failures_path = self.run_dir / "failures.jsonl"

    def _claim_result_dir(self) -> None:
        lock_path = self.run_dir / ".kkday-export.lock"
        flags = os.O_CREAT | os.O_RDWR
        if fcntl is None:
            flags |= os.O_EXCL
        try:
            lock_fd = os.open(lock_path, flags, 0o600)
            if fcntl is not None:
                fcntl.flock(lock_fd, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except (BlockingIOError, FileExistsError) as exc:
            if "lock_fd" in locals():
                os.close(lock_fd)
            raise RuntimeError(
                f"结果目录正在被另一个导出任务占用：{self.run_dir}"
            ) from exc
        self._lock_fd = lock_fd
        self._lock_path = lock_path
        self._uses_marker_lock = fcntl is None
        os.ftruncate(lock_fd, 0)
        os.write(lock_fd, f"batch_id={self.batch_id}\npid={os.getpid()}\n".encode())

    def close(self) -> None:
        if self._lock_fd is None:
            return
        try:
            if fcntl is not None and not self._uses_marker_lock:
                fcntl.flock(self._lock_fd, fcntl.LOCK_UN)
        finally:
            os.close(self._lock_fd)
            self._lock_fd = None
            if self._uses_marker_lock and self._lock_path is not None:
                self._lock_path.unlink(missing_ok=True)
            self._lock_path = None
            self._uses_marker_lock = False

    def __del__(self) -> None:
        self.close()

    def _row_values(
        self,
        row: SkuRow,
        snapshot: ProductSnapshot | None,
        *,
        include_product_fields: bool,
    ) -> list:
        product_values = [
            row.vendor,
            row.input_order + 1,
            row.product_id,
            row.product_url,
            row.product_name,
            snapshot.destination if snapshot else "",
            snapshot.rating if snapshot else None,
            snapshot.review_count if snapshot else None,
            snapshot.sold_count if snapshot else None,
            snapshot.wishlist_count if snapshot else None,
            snapshot.page_min_price if snapshot else None,
            snapshot.confirmation if snapshot else "",
            snapshot.cancellation_policy if snapshot else "",
            " | ".join(snapshot.features) if snapshot else "",
        ]
        if not include_product_fields:
            product_values = [None] * len(product_values)
        values = [
            *product_values,
            row.package_id,
            row.package_name,
            row.item_id,
            row.sku_id,
            row.usage_date.isoformat(),
            row.option_text,
            row.currency,
            row.market,
            excel_value(row.sale_price),
            excel_value(row.official_price),
            row.price_source,
            row.price_status.value,
            row.note,
            row.captured_at.isoformat(timespec="seconds"),
        ]
        return [excel_value(value) for value in values]

    def checkpoint(
        self, products: tuple[ProductResult, ...], final: bool
    ) -> Path:
        ordered_products = sorted(products, key=lambda product: product.input_item.input_order)
        product_rows: list[tuple[ProductResult, list[SkuRow]]] = []
        for product in ordered_products:
            rows = list(product.rows)
            if not self.config.include_unavailable:
                rows = [row for row in rows if row.available]
            rows.sort(key=lambda row: row.tree_order)
            product_rows.append((product, rows))
        rows = [row for _product, group in product_rows for row in group]
        failures = [
            failure for product in products for failure in product.failures
        ]
        complete_success = final and not failures
        warnings = [
            (product.input_item.url, warning)
            for product in products
            for warning in product.warnings
        ]

        workbook = openpyxl.Workbook()
        sku_sheet = workbook.active
        sku_sheet.title = "抓取结果"
        sku_sheet.append(BASE_FIELDS)
        for product, group in product_rows:
            for index, row in enumerate(group):
                sku_sheet.append(
                    self._row_values(
                        row,
                        product.snapshot,
                        include_product_fields=index == 0,
                    )
                )

        failure_sheet = workbook.create_sheet("失败商品")
        failure_sheet.append(["类型", "URL", "阶段", "消息", "可重试", "时间"])
        for failure in failures:
            failure_sheet.append(
                [
                    excel_value(value)
                    for value in [
                        failure.kind.value,
                        failure.url,
                        failure.stage,
                        failure.message,
                        "是" if failure.retryable else "否",
                        failure.captured_at.isoformat(timespec="seconds"),
                    ]
                ]
            )
        for url, warning in warnings:
            failure_sheet.append(
                [
                    excel_value(value)
                    for value in ["warning", url, "quality", warning, "否", ""]
                ]
            )
        failed_urls = {failure.url for failure in failures}
        for product, group in product_rows:
            if not group and product.input_item.url not in failed_urls:
                failure_sheet.append(
                    [
                        "no_available_sku",
                        excel_value(product.input_item.url),
                        "business_result",
                        "指定日期没有可导出的可售 SKU",
                        "否",
                        "",
                    ]
                )

        summary_sheet = workbook.create_sheet("运行汇总")
        summary_sheet.append(["字段", "值"])
        summary_sheet.append(["批次ID", self.batch_id])
        summary_sheet.append(["运行ID", self.run_id])
        summary_sheet.append(["使用日期", self.config.usage_date.isoformat()])
        summary_sheet.append(["输入商品数", self.config.requested_count or len(products)])
        summary_sheet.append(["已处理商品数", len(products)])
        summary_sheet.append(["成功商品数", sum(not product.failures for product in products)])
        summary_sheet.append(["失败商品数", sum(bool(product.failures) for product in products)])
        summary_sheet.append(["导出SKU数", len(rows)])
        summary_sheet.append(["失败数", len(failures)])
        summary_sheet.append(["Task space", self.config.ego_task_space or f"kkday-sku-price-{self.config.market}"])
        summary_sheet.append(["状态", "final" if complete_success else "partial"])

        for sheet in workbook.worksheets:
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for column in sheet.columns:
                width = min(
                    max(
                        max(len(str(cell.value or "")) for cell in column) + 2,
                        10,
                    ),
                    70,
                )
                sheet.column_dimensions[
                    get_column_letter(column[0].column)
                ].width = width

        with tempfile.NamedTemporaryFile(
            dir=self.run_dir,
            prefix=".workbook-",
            suffix=".xlsx",
            delete=False,
        ) as handle:
            temp_workbook = Path(handle.name)
        try:
            workbook.save(temp_workbook)
            target_output = self.output_path if complete_success else self.partial_output_path
            os.replace(temp_workbook, target_output)
            if complete_success:
                self.partial_output_path.unlink(missing_ok=True)
            else:
                self.output_path.unlink(missing_ok=True)
        finally:
            workbook.close()
            temp_workbook.unlink(missing_ok=True)

        summary = {
            "batch_id": self.batch_id,
            "run_id": self.run_id,
            "usage_date": self.config.usage_date.isoformat(),
            "status": "final" if complete_success else "partial",
            "requested_count": self.config.requested_count or len(products),
            "processed_count": len(products),
            "product_count": len(products),
            "row_count": len(rows),
            "failure_count": len(failures),
            "task_space": self.config.ego_task_space or f"kkday-sku-price-{self.config.market}",
            "output": str(target_output),
        }
        self._atomic_text(
            self.summary_path,
            json.dumps(summary, ensure_ascii=False, indent=2) + "\n",
        )
        if failures:
            lines = [
                json.dumps(
                    {
                        "kind": failure.kind.value,
                        "url": failure.url,
                        "stage": failure.stage,
                        "message": failure.message,
                        "retryable": failure.retryable,
                        "captured_at": failure.captured_at.isoformat(
                            timespec="seconds"
                        ),
                    },
                    ensure_ascii=False,
                )
                for failure in failures
            ]
            self._atomic_text(self.failures_path, "\n".join(lines) + "\n")
        else:
            self.failures_path.unlink(missing_ok=True)
        return target_output

    def _atomic_text(self, path: Path, content: str) -> None:
        with tempfile.NamedTemporaryFile(
            dir=path.parent,
            prefix=f".{path.name}-",
            mode="w",
            encoding="utf-8",
            delete=False,
        ) as handle:
            handle.write(content)
            temporary = Path(handle.name)
        try:
            os.replace(temporary, path)
        finally:
            temporary.unlink(missing_ok=True)
